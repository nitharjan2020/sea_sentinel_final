"""Fill a copy of SEA_Sentinel_Dataset_Inventory.xlsx with the audit results.

The template (data/templates/) is never modified; the filled copy is written to outputs/tables/.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .config import DATA_DIR, DISEASES, TABLE_DIR

TEMPLATE = DATA_DIR / "templates" / "SEA_Sentinel_Dataset_Inventory.xlsx"

# Source Register updates: Source ID -> {column header: value}
SOURCE_UPDATES = {
    "SRC001": {
        "Acquisition status": "Downloaded (2 event-list exports, 168 events)",
        "Retrieved (YYYY-MM-DD)": "2026-10-20",
        "Next action": "Event-level only; use for validation and notification delays. "
        "Bulk outbreak-level export still needed",
    },
    "SRC003": {
        "Acquisition status": "Downloaded (26,482 records; file saved as FAOSTAT.csv)",
        "Retrieved (YYYY-MM-DD)": "2026-10-21",
        "Next action": "Primary outbreak source; deduplicated within source (incl. WOAH-sourced rows)",
        "Caveats": "Includes WOAH-derived records: do not add WAHIS outbreak counts on top. "
        "Some ASF rows lack observation dates (report-date fallback, flagged)",
    },
    "SRC004": {
        "Acquisition status": "Downloaded (chickens, ducks, geese, turkeys, pigs; 2000-2024)",
        "Retrieved (YYYY-MM-DD)": "2026-10-20",
    },
    "SRC005": {
        "Acquisition status": "Downloaded (GLW chicken 2010/2015, pig 2015 rasters + census metadata)",
        "Retrieved (YYYY-MM-DD)": "2026-10-22",
    },
    "SRC006": {
        "Acquisition status": "Not downloaded: NOAA CPC daily climate (SRC011) used instead",
        "Next action": "Only if a finer or reanalysis product is needed (CDS registration required)",
    },
    "SRC007": {
        "Acquisition status": "Downloaded via API (population by sex, land area, density, GDP, trade; 2000-2025)",
        "Retrieved (YYYY-MM-DD)": "2026-10-24",
        "Next action": "Context only: annual national values are not weekly forecast predictors",
        "Caveats": "Not a disease outcome dataset; trade missing for Myanmar, Laos before 2009",
    },
}

# Sources not in the template: appended to the Source Register.
ADDED_SOURCES = {
    "SRC011": {
        "Dataset / provider": "NOAA CPC Global Unified gauge-based daily precipitation and temperature (NOAA PSL)",
        "Official / source URL": "https://psl.noaa.gov/data/gridded/data.cpc.globalprecip.html",
        "Data type": "Daily gridded climate (0.5 degree): precipitation, tmax, tmin",
        "Target disease(s)": "HPAI; ASF", "Geography": "11 SEA countries (country averages of land cells)",
        "Temporal coverage": "Daily, 2003-present", "Access format": "NetCDF subset via THREDDS NCSS (open)",
        "Acquisition status": "Downloaded (python -m sea_sentinel.download)",
        "Next action": "Tested as forecast predictors (report 8.5)", "Priority": "Medium",
        "Caveats": "National averages hide regional climate; gauge density varies", "Retrieved (YYYY-MM-DD)": "2026-10-24",
    },
    "SRC012": {
        "Dataset / provider": "Natural Earth admin-0 countries (1:50m)",
        "Official / source URL": "https://www.naturalearthdata.com/",
        "Data type": "Country boundaries", "Target disease(s)": "All (support)", "Geography": "11 SEA countries",
        "Temporal coverage": "Static", "Access format": "GeoJSON", "Acquisition status": "Downloaded",
        "Next action": "Used only to average the climate grids per country", "Priority": "Low",
        "Caveats": "Generalised 1:50m outlines", "Retrieved (YYYY-MM-DD)": "2026-10-24",
    },
    "SRC013": {
        "Dataset / provider": "NDWI (Landsat 8 Collection 2 Tier 1) and NDVI (MODIS/Terra MOD13 vegetation indices)",
        "Official / source URL": "https://lpdaac.usgs.gov/products/mod13a2v061/ ; https://www.usgs.gov/landsat-missions",
        "Data type": "Satellite water and vegetation indices", "Target disease(s)": "HPAI (wild-bird habitat); ASF",
        "Geography": "SEA", "Temporal coverage": "16-day composites (MODIS), 16-day revisit (Landsat 8)",
        "Access format": "Google Earth Engine or NASA AppEEARS", "Acquisition status": "Not downloaded",
        "Next action": "Needs an Earth Engine or NASA Earthdata account; see report 10.3", "Priority": "Medium",
        "Caveats": "Cloud cover; aggregate to admin1 x month", "Retrieved (YYYY-MM-DD)": "",
    },
}

CHECKLIST_EMPRES = {
    "Event identity": "Yes: Event.ID unique per row; duplicates handled by location/date rules",
    "Disease": "Yes: LPAI and non-HPAI serotypes excluded",
    "Species": "Yes: domestic poultry / domestic swine only (rabies dropped from scope)",
    "Occurrence date": "Partial: observation date missing for some ASF rows (report-date fallback)",
    "Report date": "Yes: >99.9% present",
    "Location": "Yes: admin1, locality and coordinates on every row",
    "Outcome": "Yes: outbreaks only (no animal case counts used)",
    "Coverage": "Yes: zeros only within first-last report window",
    "Provenance": "Yes: data/raw/MANIFEST.csv (SHA-256, retrieval date)",
    "Permissions": "Yes: FAO databases CC BY 4.0 (attribution); see DATA_LICENSES.md",
    "Validation": "See Coverage Matrix 'Forecast eligibility'",
}


def _header_index(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _cell_value(v):
    if isinstance(v, pd.Timestamp):
        return v.date() if not pd.isna(v) else "None"
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return v


def write_inventory(audit: pd.DataFrame, manifest: pd.DataFrame, out_dir: Path = TABLE_DIR) -> Path:
    wb = load_workbook(TEMPLATE)
    reg, cov, chk = wb["Source Register"], wb["Coverage Matrix"], wb["Quality Checklist"]

    cols = _header_index(reg)
    for row in reg.iter_rows(min_row=2):
        sid = row[cols["Source ID"] - 1].value
        for header, value in SOURCE_UPDATES.get(sid, {}).items():
            if header in cols:
                row[cols[header] - 1].value = value

    present = {row[cols["Source ID"] - 1].value for row in reg.iter_rows(min_row=2)}
    for sid, values in ADDED_SOURCES.items():
        if sid not in present:
            reg.append([sid if h == "Source ID" else values.get(h, "") for h in cols])

    cols = _header_index(cov)
    lookup = audit.set_index(["ISO3", "Disease"])
    out_of_scope = []  # template rows for diseases dropped from scope (rabies)
    for row in cov.iter_rows(min_row=2):
        key = (row[cols["ISO3"] - 1].value, row[cols["Disease"] - 1].value)
        if key not in lookup.index:
            if key[1] and key[1] not in DISEASES:
                out_of_scope.append(row[0].row)
            continue
        rec = lookup.loc[key]
        for header, col in cols.items():
            if header in rec.index and header not in ("ISO3", "Disease"):
                row[col - 1].value = _cell_value(rec[header])
    for r in reversed(out_of_scope):
        cov.delete_rows(r)

    cols = _header_index(chk)
    for row in chk.iter_rows(min_row=2):
        cat = row[cols["Category"] - 1].value
        if cat in CHECKLIST_EMPRES:
            row[cols["Verified?"] - 1].value = "EMPRES-i: " + CHECKLIST_EMPRES[cat].split(":")[0]
            row[cols["Notes"] - 1].value = CHECKLIST_EMPRES[cat]

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "SEA_Sentinel_Dataset_Inventory_assessed.xlsx"
    wb.save(out)
    return out
